import io
import json
import sys
from pathlib import Path
from unittest.mock import AsyncMock

import httpx
import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).parent))
import public_app as public

RealClient = httpx.AsyncClient

def document_bytes():
    document = Document()
    document.add_heading("Fictional Test Candidate", 0)
    document.add_paragraph("Operations administrator with experience maintaining records, coordinating schedules, and supporting customers.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()

@pytest.mark.asyncio
async def test_public_has_no_login_or_shared_resume_routes():
    async with RealClient(transport=httpx.ASGITransport(app=public.app), base_url="http://test") as client:
        home = await client.get("/")
        assert home.status_code == 200
        assert "No account needed" in home.text
        assert home.headers["cache-control"] == "private, no-store"
        assert (await client.get("/api/v1/resumes/list")).status_code == 404
        assert (await client.get("/settings")).status_code == 404
        assert (await client.get("/api/health")).json()["mode"] == "public-stateless"

@pytest.mark.asyncio
async def test_keys_and_outputs_are_request_scoped(monkeypatch):
    monkeypatch.setattr(public, "parse_document", AsyncMock(return_value="Fictional Test Candidate has substantial experience coordinating customer records and administrative schedules."))
    seen = []
    class ProviderClient:
        def __init__(self, **kwargs):
            pass
        async def __aenter__(self):
            return self
        async def __aexit__(self, *args):
            pass
        async def post(self, url, headers, json):
            seen.append((url, headers["Authorization"], json))
            name = "FIRST PRIVATE DRAFT" if headers["Authorization"].endswith("key-first-test") else "SECOND PRIVATE DRAFT"
            result = {"resume_markdown": "# " + name + "\n" + "Fictional experience supporting customers and maintaining administrative records. " * 3, "changes": ["Reordered existing experience"], "warnings": []}
            return httpx.Response(200, json={"choices": [{"finish_reason": "stop", "message": {"content": __import__("json").dumps(result)}}]})
    monkeypatch.setattr(public.httpx, "AsyncClient", ProviderClient)
    async with RealClient(transport=httpx.ASGITransport(app=public.app), base_url="http://test") as client:
        outputs = []
        for key in ("key-first-test", "key-second-test"):
            response = await client.post("/api/tailor", files={"file": ("test.docx", document_bytes())},
                data={"job_description": "An operations role maintaining accurate records and coordinating administrative customer schedules. " * 2,
                      "provider": "groq", "model": "openai/gpt-oss-120b", "api_key": key})
            assert response.status_code == 200, response.text
            outputs.append(response.json())
        assert "Fictional Test Candidate" in outputs[0]["source_markdown"]
        assert "FIRST PRIVATE DRAFT" not in outputs[0]["source_markdown"]
        assert "FIRST PRIVATE DRAFT" in outputs[0]["resume_markdown"]
        assert "FIRST PRIVATE DRAFT" not in outputs[1]["resume_markdown"]
        assert seen[0][1] != seen[1][1]
        assert "key-first-test" not in json.dumps(outputs)
        assert seen[0][0] == public.PROVIDERS["groq"]

@pytest.mark.asyncio
async def test_docx_export_and_html_escaping():
    async with RealClient(transport=httpx.ASGITransport(app=public.app), base_url="http://test") as client:
        response = await client.post("/api/export", json={"markdown": "# Fictional Candidate\n## Experience\n- Coordinated records", "format": "docx"})
        assert response.status_code == 200
        document = Document(io.BytesIO(response.content))
        assert any(p.text == "Fictional Candidate" for p in document.paragraphs)
    rendered = public.resume_html("# <script>alert(1)</script>\n![image](http://internal)")
    assert "<script>" not in rendered
    assert "&lt;script&gt;" in rendered
    assert "<img" not in rendered

@pytest.mark.asyncio
async def test_validation_prevents_provider_calls():
    async with RealClient(transport=httpx.ASGITransport(app=public.app), base_url="http://test") as client:
        bad = await client.post("/api/tailor", files={"file": ("payload.exe", b"invalid")}, data={
            "job_description": "A sufficiently long fictional job description maintaining customer records and schedules." * 2,
            "provider": "groq", "model": "openai/gpt-oss-120b", "api_key": "key-test-only"})
        assert bad.status_code == 400
        assert (await client.post("/api/export", json={"markdown": "", "format": "pdf"})).status_code == 422
